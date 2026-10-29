
from pathlib import Path
from datetime import datetime, date
import csv
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
PROCESSED = BASE / "processed"
LOGS = BASE / "logs"

PROCESSED.mkdir(exist_ok=True)
LOGS.mkdir(exist_ok=True)

issues = []


def record_issue(source, row, issue):
    issues.append([source, row, issue])

def valid_date(value):
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    return None

def save_csv(path, headers, rows):
    with open(path, "w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)


# 1. Read original RBA data
rba = load_workbook(BASE / "RBA raw data.xlsx", read_only=True, data_only=True)
rba_sheet = rba["Data"]

rba_rows = []
rba_dates = set()

# RBA observations begin at row 12
for row_number, row in enumerate(
    rba_sheet.iter_rows(min_row=12, max_col=2, values_only=True),
    start=12
):
    observation_date, cash_rate = row
    d = valid_date(observation_date)

    if d is None:
        record_issue("RBA", row_number, "Missing or invalid date")
        continue

    if d in rba_dates:
        record_issue("RBA", row_number, "Duplicate date")
        continue

    if not isinstance(cash_rate, (int, float)) or isinstance(cash_rate, bool):
        record_issue("RBA", row_number, "Missing or invalid cash rate")
        continue

    rba_dates.add(d)
    rba_rows.append([d.isoformat(), cash_rate])

rba.close()
rba_rows.sort(key=lambda row: row[0])


# 2. Read original ASX 200 data
asx = load_workbook(BASE / "AXS 200.xlsx", read_only=True, data_only=True)

# Find the raw sheet
asx_sheet = next(
    sheet for sheet in asx.worksheets
    if sheet.title.strip().lower() == "raw data"
)

asx_rows = []
asx_dates = set()

for row_number, row in enumerate(
    asx_sheet.iter_rows(min_row=2, max_col=7, values_only=True),
    start=2
):
    observation_date, open_price, high, low, close, adj_close, volume = row
    d = valid_date(observation_date)

    if d is None:
        record_issue("ASX 200", row_number, "Missing or invalid date")
        continue

    if d in asx_dates:
        record_issue("ASX 200", row_number, "Duplicate date")
        continue

    prices = [open_price, high, low, close, adj_close]
    if any(
        not isinstance(value, (int, float)) or isinstance(value, bool)
        for value in prices
    ):
        record_issue("ASX 200", row_number, "Missing or invalid price")
        continue

    if not isinstance(volume, (int, float)) or isinstance(volume, bool):
        record_issue("ASX 200", row_number, "Missing or invalid volume")
        volume = ""

    asx_dates.add(d)
    asx_rows.append([
        d.isoformat(), open_price, high, low, close, adj_close, volume
    ])

asx.close()
asx_rows.sort(key=lambda row: row[0])


# 3. Save cleaned data as CSV files
save_csv(
    PROCESSED / "rba_clean.csv",
    ["Date", "Cash Rate Target (%)"],
    rba_rows
)

save_csv(
    PROCESSED / "asx200_clean.csv",
    ["Date", "Open", "High", "Low", "Close", "Adj Close", "Volume"],
    asx_rows
)

# 4. Save data quality log
save_csv(
    LOGS / "data_quality_log.csv",
    ["Source", "Excel Row", "Issue"],
    issues
)

print("Pipeline completed!")
print(f"RBA valid records: {len(rba_rows)}")
print(f"ASX 200 valid records: {len(asx_rows)}")
print(f"Quality issues logged: {len(issues)}")
print("Outputs saved in processed/ and logs/")
